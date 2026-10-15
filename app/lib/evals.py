"""
Evals: the checks that prove the dashboard's numbers are right (spec, "Done Means").

Each check returns a dict: item (the spec's Done Means number), name, passed, public, private.
`public` never contains amounts (it goes to Outputs/, which syncs to OneDrive); `private` has the figures
and goes to FinanceData\\evals\\ only.

Checks compare against something independent of the code under test: the raw Excel read with openpyxl (not
our import code), direct SQL, the platform's own interest figures, and the replay's truth table.
"""
import os
import re
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.lib import config as cfg
from app.lib import db, fmt, replay
from app.lib import interest as it
from app.lib import metrics as mx

OPENING = ("2026-09-28", 1_234_567.89)  # spec: net worth on the opening snapshot
DEPOSIT_TOLERANCE = 1.00                 # pesos
VIEWS = cfg.PROJECT_DIR / "app" / "views"
RAW_TYPES = {"Exp.": "expense", "Income": "income", "Transfer-Out": "transfer_out", "Transfer-In": "transfer_in"}


def _check(item, name, passed, public, private=None):
    return {"item": item, "name": name, "passed": bool(passed), "public": public, "private": private or public}


# --- 1. Money Manager import ------------------------------------------------------

def _raw_rows(path):
    """(timestamp, type, category, amount) for every row, read straight from the Excel with openpyxl."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb.active.iter_rows(values_only=True)
        header = list(next(rows))
        i = {n: header.index(n) for n in ("Period", "Category", "Income/Expense", "Amount")}
        out = []
        for r in rows:
            if r[i["Period"]] is None:
                continue
            when = r[i["Period"]]
            when = when if isinstance(when, datetime) else pd.to_datetime(when).to_pydatetime()
            out.append((when, RAW_TYPES.get(r[i["Income/Expense"]], r[i["Income/Expense"]]),
                        str(r[i["Category"]] or "").strip(), round(float(r[i["Amount"]]), 2)))
        return out
    finally:
        wb.close()


_NAME = re.compile(r"(\d{4}-\d{2}-\d{2})\s*~\s*(\d{4}-)?(\d{2}-\d{2})")


def _owned(path, rows):
    """The dates an export speaks for (spec rule, decided 2026-09-29), written separately from the import code:
    the range in the file name, or up to its last row if saved before that range ended; no name: its rows."""
    days = sorted(r[0].date() for r in rows)
    if not days:
        return None
    m = _NAME.search(Path(path).name)
    if not m:
        return days[0], days[-1]
    start = date.fromisoformat(m.group(1))
    end = date.fromisoformat(f"{m.group(2)[:4] if m.group(2) else start.year}-{m.group(3)}")
    if end < start:
        end = end.replace(year=end.year + 1)
    if end > date.fromtimestamp(os.path.getmtime(path)):
        end = days[-1]
    return start, end


def _combined(paths):
    """Rows the dashboard should hold: newest file first, each keeping the dates no newer file owns."""
    files = [(os.path.getmtime(p), Path(p).name, p, _raw_rows(p)) for p in paths]
    owned, out, replaced = [], [], 0
    for _, _, p, rows in sorted(files, key=lambda f: (f[0], (_owned(f[2], f[3]) or (None, date.min))[1], f[1]),
                                reverse=True):
        for r in rows:
            if any(a <= r[0].date() <= b for a, b in owned):
                replaced += 1
            else:
                out.append(r)
        span = _owned(p, rows)
        if span:
            owned.append(span)
    return out, len(files), replaced


def mm_import(export_paths, db_path):
    raw = {}
    rows, n_files, replaced = _combined(export_paths)
    for when, kind, cat, amount in rows:
        k = (when.strftime("%Y-%m"), kind, cat)
        c, s = raw.get(k, (0, 0.0))
        raw[k] = (c + 1, s + amount)
    got = {(m, t, cat): (n, s) for m, t, cat, n, s in db.query(
        db_path, "SELECT month, type, category, COUNT(*), SUM(amount) FROM transactions GROUP BY 1, 2, 3")}
    diffs = []
    for k in sorted(set(raw) | set(got)):
        (rc, rs), (gc, gs) = raw.get(k, (0, 0.0)), got.get(k, (0, 0.0))
        if rc != gc or abs(rs - gs) > 0.005:
            diffs.append(f"{' / '.join(k)}: Excel {rc} filas {rs:,.2f}, base {gc} filas {gs:,.2f}")
    notes = db.read_optional(db_path, "import_warnings", ["message"])["message"].tolist()
    public = (f"{len(rows)} rows from {n_files} files ({replaced} replaced by a newer export of the same dates), "
              f"{len(raw)} month × type × category groups; {len(diffs)} differ; "
              f"{len(notes)} import warnings (holes between exports, old files saved again)")
    return _check(1, "Money Manager import matches the Excel", not diffs and not notes, public,
                  public + ("" if not diffs else ": " + "; ".join(diffs[:20])))


def rebuild(export_paths, snapshots_csv, adjustments_csv, settings, workdir):
    workdir = Path(workdir)
    workdir.mkdir(parents=True, exist_ok=True)
    a, b, c = workdir / "build-1.db", workdir / "build-2.db", workdir / "build-twice-listed.db"
    ra = db.build_database(a, export_paths, snapshots_csv, adjustments_csv, settings)
    rb = db.build_database(b, export_paths, snapshots_csv, adjustments_csv, settings)
    rc_ = db.build_database(c, list(export_paths) * 2, snapshots_csv, adjustments_csv, settings)  # every file twice
    same = db.fingerprint(a) == db.fingerprint(b)
    counts = {k: (ra[k], rb[k]) for k in ("transactions", "snapshot_items", "adjustments")}
    ok = same and all(x == y for x, y in counts.values()) and rc_["transactions"] == ra["transactions"]
    public = (f"built twice: {'identical' if same else 'DIFFERENT'} contents; rows "
              + ", ".join(f"{k} {x} then {y}" for k, (x, y) in counts.items())
              + f"; with every export listed twice: {rc_['transactions']} transactions")
    return _check(1, "Loading twice adds nothing", ok, public)


# --- 2. Net worth ------------------------------------------------------------------

def opening_net_worth(db_path, as_of=OPENING[0], expected=OPENING[1]):
    series = mx.net_worth_series(db.read_table(db_path, "snapshot_items")).set_index("as_of")["net_worth"]
    got = series.get(as_of)
    ok = got is not None and abs(got - expected) < 0.005
    return _check(2, "Opening net worth", ok, f"{as_of}: {'matches the spec' if ok else 'DOES NOT match the spec'}",
                  f"{as_of}: base {got if got is None else fmt.exact(got)} vs. spec {fmt.exact(expected)}")


@contextmanager
def _database(db_path):
    """Point the dashboard at another database for the length of the block."""
    old = os.environ.get("FINANCE_DB")
    os.environ["FINANCE_DB"] = str(db_path)
    try:
        yield
    finally:
        if old is None:
            os.environ.pop("FINANCE_DB", None)
        else:
            os.environ["FINANCE_DB"] = old


def _open(view, db_path):
    from streamlit.testing.v1 import AppTest
    with _database(db_path):
        at = AppTest.from_file(str(VIEWS / view), default_timeout=60).run()
    if at.exception:
        raise RuntimeError(f"{view}: {[e.value for e in at.exception]}")
    return at


def _card(at, label):
    m = next((m for m in at.metric if m.label == label), None)
    return None if m is None else m.help


def dashboard_net_worth(db_path, expected=None):
    """The Resumen "Patrimonio" card shows the latest net worth in the database (or `expected`)."""
    series = mx.net_worth_series(db.read_table(db_path, "snapshot_items"))
    as_of, value = series["as_of"].iloc[-1], float(series["net_worth"].iloc[-1])
    want = f"{fmt.exact(value if expected is None else expected)} al {fmt.day(as_of)}"
    got = _card(_open("resumen.py", db_path), "Patrimonio")
    ok = got == want
    return _check(2, "Net worth card on the dashboard", ok,
                  f"Resumen card for {as_of}: {'matches' if ok else 'DOES NOT match'} the "
                  f"{'database' if expected is None else 'expected figure'}",
                  f"Resumen card: {got!r} vs. expected {want!r}")


# --- 3. Fixed-yield math -------------------------------------------------------------

def deposit_math(snapshots_csv, settings):
    s = pd.read_csv(snapshots_csv, encoding="utf-8-sig", dtype={"as_of": str, "start_date": str, "end_date": str})
    all_deps = s[s["item_type"] == "deposit"]
    deps = all_deps[all_deps["interest_total"].notna()]
    skipped = len(all_deps) - len(deps)
    isr = settings["tax"]["isr_rate"]
    bad, lines = [], []
    for r in deps.itertuples():
        calc = it.deposit_interest(r, isr)
        diff = round(calc - r.interest_total, 2)
        lines.append(f"{r.as_of} {r.platform} {r.name}: mostrado {r.interest_total:,.2f}, calculado {calc:,.2f}")
        if abs(diff) > DEPOSIT_TOLERANCE:
            bad.append(lines[-1])
    public = (f"{len(deps) - len(bad)} of {len(deps)} deposit interest figures within 1 peso of the platform's; "
              f"{skipped} deposits had no figure to compare")
    return _check(3, "Fixed-yield math", deps.shape[0] > 0 and not bad, public,
                  public + ": " + "; ".join(lines))


# --- 4. Filters ------------------------------------------------------------------------

def _sql(db_path, kind, period, accounts=None, categories=None):
    sql = "SELECT COUNT(*), COALESCE(SUM(amount), 0) FROM transactions WHERE type = ? AND date BETWEEN ? AND ?"
    params = [kind, str(period[0]), str(period[1])]
    for col, values in (("account", accounts), ("category", categories)):
        if values:
            sql += f" AND {col} IN ({','.join('?' * len(values))})"
            params += list(values)
    n, total = db.query(db_path, sql, params)[0]
    return n, round(total, 2)


def filters(db_path, pages=True):
    """Totals for a grid of filters, through the dashboard's own code, vs. a direct SQL query."""
    from app.lib import ui  # quick periods live with the screens
    tx = db.read_table(db_path, "transactions")
    snaps = db.read_table(db_path, "snapshot_items")
    first = date.fromisoformat(min(tx["date"].min(), snaps["as_of"].astype(str).min()))
    last = date.fromisoformat(max(tx["date"].max(), snaps["as_of"].astype(str).max()))
    periods = {p: ui.preset_range(p, first, last) for p in ui.PRESETS}
    periods = {p: r for p, r in periods.items() if r}

    def top(kind, col, n):
        t = tx[tx["type"] == kind]
        return list(t.groupby(col)["amount"].sum().sort_values(ascending=False).index[:n])

    combos = []
    for kind in ("expense", "income"):
        accs = [None, top(kind, "account", 1), top(kind, "account", 2)]
        cats = [None, top(kind, "category", 1)]
        combos += [(kind, p, a, c) for p in periods for a in accs for c in cats]

    bad = []
    for kind, p, a, c in combos:
        f = mx.filter_tx(tx, dates=periods[p], accounts=a, categories=c)
        f = f[f["type"] == kind]
        got = (len(f), round(float(f["amount"].sum()), 2))
        want = _sql(db_path, kind, periods[p], a, c)
        if got[0] != want[0] or abs(got[1] - want[1]) > 0.005:
            bad.append(f"{kind} {p} {a} {c}: código {got}, SQL {want}")

    page_runs = 0
    if pages:
        from streamlit.testing.v1 import AppTest
        screens = [("gastos.py", "expense", "w_accounts_exp", "w_categories_exp"),
                   ("ingresos.py", "income", "w_accounts_inc", "w_categories_inc")]
        for view, kind, acc_key, cat_key in screens:
            grid = [(p, None, None) for p in periods]
            grid += [("12 meses", top(kind, "account", 1), None), ("12 meses", None, top(kind, "category", 1)),
                     ("Todo", top(kind, "account", 2), top(kind, "category", 1))]
            for p, a, c in grid:
                if p not in periods:
                    continue
                with _database(db_path):
                    at = AppTest.from_file(str(VIEWS / view), default_timeout=60).run()
                    at.date_input[0].set_value(periods[p]).run()
                    if a:
                        at.multiselect(key=acc_key).set_value(a).run()
                    if c:
                        at.multiselect(key=cat_key).set_value(c).run()
                page_runs += 1
                n, total = _sql(db_path, kind, periods[p], a, c)
                got = _card(at, "Total")
                if at.exception or (got != fmt.exact(total) and not (n == 0 and got is None)):
                    bad.append(f"{view} {p} {a} {c}: tarjeta {got!r}, SQL {fmt.exact(total)} "
                               f"{[e.value for e in at.exception] if at.exception else ''}")
                if kind == "income" and got is not None:
                    # Net income and savings rate ignore the source filter: they need every expense too
                    inc = _sql(db_path, "income", periods[p], a)[1]
                    exp_ = _sql(db_path, "expense", periods[p], a)[1]
                    card = next((m for m in at.metric if m.label == "Ingreso neto"), None)
                    want = (fmt.exact(inc - exp_), f"Tasa de ahorro {fmt.pct((inc - exp_) / inc if inc else None)}")
                    have = (card.help, card.delta) if card else None
                    if have != want:
                        bad.append(f"{view} {p} {a} {c}: ingreso neto {have!r}, SQL {want!r}")

    public = (f"{len(combos)} filter combinations in code and {page_runs} on the Gastos/Ingresos screens "
              f"(Total, net income and savings rate cards) vs. direct SQL; {len(bad)} differ")
    return _check(4, "Filters match the database", not bad, public, public + ("" if not bad else ": " + "; ".join(bad)))


# --- 5. Replay -------------------------------------------------------------------------

def replay_check(workdir, baseline_csv, export_paths, settings):
    r = replay.run(workdir, baseline_csv, export_paths, settings)
    card = dashboard_net_worth(r.db, expected=float(r.truth["net_worth"].iloc[-1]))
    checks = r.checks + [{"month": r.months[-1], "name": "dashboard card", "passed": card["passed"],
                          "detail": card["private"]}]
    failed = [f"{c['month']} {c['name']}" for c in checks if not c["passed"]]
    planted = {"missing screenshot refused": "missing screenshot",
               "matured deposit still listed warned": "matured deposit still listed",
               "duplicate flagged": "duplicate row", "logging mistake flagged": "transfer logged with the wrong amount",
               "stock sale flagged": "stock sale", "new category, account and refund loaded":
               "new category, new account and refund",
               "re-sent month not double counted": "month re-sent with a tag removed",
               "vanished rows warned": "row deleted in a re-sent month"}
    # "vanished rows warned" runs every month (none expected); the planted deletion is in month 3
    status = [f"{label} ({c['month'][:7]}: {'caught' if c['passed'] else 'MISSED'})"
              for name, label in planted.items() for c in checks if c["name"] == name
              and (name != "vanished rows warned" or c["month"] == r.months[replay.SALE_AT])]
    public = (f"{len(checks) - len(failed)} of {len(checks)} replay checks pass over {len(r.months)} months "
              f"({r.months[0]} to {r.months[-1]}), incl. net worth each month vs. the truth table. "
              f"Planted errors: {'; '.join(status)}." + (f" Failed: {', '.join(failed)}." if failed else ""))
    private = public + "\n" + "\n".join(
        f"  - {c['month']} {'OK ' if c['passed'] else 'FAIL'} {c['name']}: {c['detail']}" for c in checks)
    return _check(5, "Replay of 3 months with planted errors", not failed, public, private)


# --- All checks and the report ----------------------------------------------------------

def run_all(settings, workdir):
    """Every check on the real files. Scratch builds go in `workdir` (a new, empty folder)."""
    data_dir = Path(settings["paths"]["data_dir"])
    exports = cfg.find_exports(data_dir)  # skips the "_" scratch folders (e.g. _replay)
    snaps_csv, adj = cfg.data_path(settings, "snapshots_csv"), cfg.data_path(settings, "adjustments_csv")
    real_db = cfg.data_path(settings, "database")
    workdir = Path(workdir)
    return [
        mm_import(exports, real_db),
        rebuild(exports, snaps_csv, adj, settings, workdir / "rebuild"),
        opening_net_worth(real_db),
        dashboard_net_worth(real_db),
        deposit_math(snaps_csv, settings),
        filters(real_db),
        replay_check(workdir / "replay", snaps_csv, exports, settings),
    ]


NOT_HERE = [
    ("6", "Screenshot reading accuracy", "Moved to step 7: measured once the reading skill exists."),
    ("7", "Independent verifier", "Runs after this report (step 6d)."),
    ("8", "Real monthly update under 20 minutes", "Step 8, the first live run."),
]


def report(checks, with_amounts, when=None):
    when = when or datetime.now().strftime("%Y-%m-%d %H:%M")
    passed = sum(c["passed"] for c in checks)
    lines = [
        "# Personal Finance Dashboard: Eval Report",
        "",
        f"Run {when}. **{passed} of {len(checks)} checks pass.**"
        + ("" if with_amounts else " This copy has no amounts; the full one is in FinanceData\\evals."),
        "",
        "Numbers from the spec's \"Done Means\" list. Re-run any time with `tools/run_evals.py`.",
        "",
        "| # | Check | Result | Detail |",
        "|---|---|---|---|",
    ]
    for c in checks:
        detail = (c["private"] if with_amounts else c["public"]).split("\n")[0].replace("|", "/")
        lines.append(f"| {c['item']} | {c['name']} | {'PASS' if c['passed'] else '**FAIL**'} | {detail} |")
    for item, name, why in NOT_HERE:
        lines.append(f"| {item} | {name} | not in this run | {why} |")
    if with_amounts:
        lines += ["", "## Details", ""]
        for c in checks:
            lines += [f"### {c['item']}. {c['name']}", "", c["private"], ""]
    return "\n".join(lines) + "\n"
